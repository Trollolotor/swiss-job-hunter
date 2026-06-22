import json
from io import BytesIO

import pytest
from pydantic import BaseModel

from cv_sections import (
    CVSections,
    compile_cv_text,
    normalize_keywords,
    parse_legacy_cv,
)


def test_keywords_split_and_dedupe():
    assert normalize_keywords(["AI engineer, LLM engineer", "MLOps; ai ENGINEER"]) == [
        "AI engineer", "LLM engineer", "MLOps",
    ]


def test_legacy_cv_sections_preserve_unclassified_text():
    sections = parse_legacy_cv("Jane Doe\n\nExperience\nBuilt systems\n\nEducation\nETH")
    assert sections["experience"] == "Built systems"
    assert sections["education"] == "ETH"
    assert "Jane Doe" in sections["additional"]
    compiled = compile_cv_text(sections)
    assert "## Experience" in compiled
    assert "## Additional" in compiled


def test_empty_sections_are_not_compiled():
    assert compile_cv_text(CVSections(about="Hello")) == "## About\nHello"


def test_txt_and_docx_extraction():
    from api_ext import _extract_cv_text
    from docx import Document

    text, kind, warnings = _extract_cv_text("cv.txt", "text/plain", b"Candidate CV")
    assert (text, kind, warnings) == ("Candidate CV", "txt", [])

    document = Document()
    document.add_heading("Experience", level=1)
    document.add_paragraph("Built reliable systems")
    buffer = BytesIO()
    document.save(buffer)
    text, kind, _ = _extract_cv_text(
        "cv.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        buffer.getvalue(),
    )
    assert kind == "docx"
    assert "Built reliable systems" in text


def test_scanned_pdf_is_rejected():
    from api_ext import _extract_cv_text
    from fastapi import HTTPException
    from reportlab.pdfgen.canvas import Canvas

    buffer = BytesIO()
    canvas = Canvas(buffer)
    canvas.showPage()
    canvas.save()
    with pytest.raises(HTTPException, match="text layer"):
        _extract_cv_text("scan.pdf", "application/pdf", buffer.getvalue())


@pytest.mark.asyncio
async def test_structured_response_repairs_with_same_model(monkeypatch):
    from llm import structured

    calls = []

    class Payload(BaseModel):
        value: str

    async def fake_call_llm(**kwargs):
        calls.append(kwargs)
        if len(calls) == 1:
            return "not json", "openai/test"
        return json.dumps({"value": "fixed"}), "openai/test"

    monkeypatch.setattr(structured, "call_llm", fake_call_llm)
    payload, model = await structured.call_structured(
        user="x", system="y", max_tokens=20, operation="cv_parsing", schema=Payload
    )
    assert payload.value == "fixed"
    assert model == "openai/test"
    assert calls[1]["model"] == "openai/test"


@pytest.mark.asyncio
async def test_tailoring_returns_full_sections(monkeypatch):
    from llm import cv_tailor
    from llm.cv_tailor import TailoredCVPayload

    async def fake_structured(**_kwargs):
        return TailoredCVPayload(
            missing_keywords=["kubernetes"],
            suggestions=[],
            tailored_cv_sections=CVSections(
                about="Tailored summary",
                experience="Existing factual experience",
            ),
        ), "openai/test"

    monkeypatch.setattr(cv_tailor, "call_structured", fake_structured)
    job = type("Job", (), {
        "title": "Platform Engineer",
        "company": "Example",
        "description": "Requirements\nPython and Kubernetes\n" * 20,
    })()
    result = await cv_tailor.tailor_cv(
        job,
        "Original",
        cv_sections={"about": "Original", "experience": "Existing factual experience"},
    )
    assert result["tailored_cv_sections"]["about"] == "Tailored summary"
    assert result["tailored_cv_sections"]["experience"] == "Existing factual experience"
    assert result["model"] == "openai/test"
